"""
测试用例智能生成服务 - 基于 AutoGen 0.7.5
使用多模态大模型分析图片，生成专业测试用例
"""
import os
import json
import asyncio
from typing import AsyncGenerator, Optional, List, Dict, Any
from io import BytesIO
from datetime import datetime

from autogen_agentchat.agents import AssistantAgent
from autogen_agentchat.messages import TextMessage, MultiModalMessage
from autogen_agentchat.base import TaskResult
from autogen_core.models import ModelFamily
from autogen_ext.models.openai import OpenAIChatCompletionClient
from autogen_core import Image as AGImage
from dotenv import load_dotenv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

load_dotenv()


class TestCaseService:
    """测试用例生成服务"""
    
    def __init__(self) -> None:
        # 多模态模型客户端（用于分析图片）
        self._vision_client = OpenAIChatCompletionClient(
            model=os.getenv("VISION_MODEL", "Qwen/Qwen2-VL-72B-Instruct"),
            base_url=os.getenv("OPENAI_BASE_URL", "https://api.siliconflow.cn/v1"),
            api_key=os.getenv("OPENAI_API_KEY", ""),
            model_info={
                "vision": True,
                "function_calling": True,
                "json_output": True,
                "family": ModelFamily.UNKNOWN,
                "structured_output": True,
                "multiple_system_messages": True,
            },
        )
        
        # 文本模型客户端（用于生成测试用例）
        self._text_client = OpenAIChatCompletionClient(
            model=os.getenv("TEXT_MODEL", "Qwen/Qwen2.5-72B-Instruct"),
            base_url=os.getenv("OPENAI_BASE_URL", "https://api.siliconflow.cn/v1"),
            api_key=os.getenv("OPENAI_API_KEY", ""),
            model_info={
                "vision": False,
                "function_calling": True,
                "json_output": True,
                "family": ModelFamily.UNKNOWN,
                "structured_output": True,
                "multiple_system_messages": True,
            },
        )
        
        # 图片分析 Agent
        self._vision_agent = AssistantAgent(
            name="vision_analyzer",
            model_client=self._vision_client,
            system_message="""你是一位专业的测试需求分析专家，擅长分析产品需求文档、思维导图、流程图和界面设计图。

你的任务是：
1. 仔细分析用户上传的图片内容
2. 识别图中的功能模块、业务流程、界面元素
3. 提取测试需求点和验收标准
4. 输出结构化的测试需求分析

输出格式要求：
- 功能模块清单
- 业务流程描述
- 关键测试点
- 潜在的边界条件和异常场景""",
            model_client_stream=True,
        )
        
        # 测试用例生成 Agent
        self._testcase_agent = AssistantAgent(
            name="testcase_generator",
            model_client=self._text_client,
            system_message="""你是一位资深测试工程师，精通测试用例设计方法（等价类划分、边界值分析、场景法、决策表等）。

你的任务是根据需求分析生成专业、全面的测试用例。

测试用例必须包含以下字段：
1. 用例编号（TC-001, TC-002...）
2. 用例标题（简洁描述测试目的）
3. 所属模块
4. 前置条件
5. 测试步骤（详细、可执行）
6. 预期结果（明确、可验证）
7. 优先级（P0-阻塞/P1-高/P2-中/P3-低）
8. 测试类型（功能/性能/兼容性/安全性/易用性）

生成原则：
- 覆盖所有功能路径（正常+异常）
- 考虑边界值和等价类
- 包含正向和反向测试
- 考虑数据组合和场景覆盖
- 输出格式为结构化JSON""",
            model_client_stream=True,
        )

    async def close(self) -> None:
        await self._vision_client.close()
        await self._text_client.close()

    async def analyze_image_stream(
        self, 
        image_data: bytes, 
        context: str = "",
        requirements: str = ""
    ) -> AsyncGenerator[str, None]:
        """
        流式分析图片内容
        
        Args:
            image_data: 图片二进制数据
            context: 上下文背景信息
            requirements: 用户具体要求
        """
        try:
            # 创建图片对象
            ag_image = AGImage.from_bytes(image_data)
            
            # 构建提示词
            prompt = self._build_vision_prompt(context, requirements)
            
            # 创建多模态消息
            multi_modal_message = MultiModalMessage(
                content=[prompt, ag_image],
                source="user"
            )
            
            # 流式分析
            stream = self._vision_agent.run_stream(task=multi_modal_message)
            previous_content = ""
            
            async for event in stream:
                current_content = self._extract_text(event)
                if current_content is None:
                    continue
                    
                # 计算增量内容
                if len(current_content) > len(previous_content):
                    delta = current_content[len(previous_content):]
                    previous_content = current_content
                    if delta:
                        yield json.dumps({
                            "type": "analysis",
                            "content": delta,
                            "stage": "image_analysis"
                        }, ensure_ascii=False)
            
            # 保存分析结果供后续使用
            self._last_analysis = previous_content
            
        except Exception as e:
            yield json.dumps({
                "type": "error",
                "error": f"图片分析失败: {str(e)}"
            }, ensure_ascii=False)

    async def generate_testcases_stream(
        self,
        analysis_result: str,
        context: str = "",
        requirements: str = ""
    ) -> AsyncGenerator[str, None]:
        """
        流式生成测试用例
        
        Args:
            analysis_result: 图片分析结果
            context: 上下文背景信息
            requirements: 用户具体要求
        """
        try:
            # 构建测试用例生成提示词
            prompt = self._build_testcase_prompt(analysis_result, context, requirements)
            
            # 流式生成
            stream = self._testcase_agent.run_stream(task=prompt)
            previous_content = ""
            
            async for event in stream:
                current_content = self._extract_text(event)
                if current_content is None:
                    continue
                    
                # 计算增量内容
                if len(current_content) > len(previous_content):
                    delta = current_content[len(previous_content):]
                    previous_content = current_content
                    if delta:
                        yield json.dumps({
                            "type": "testcase",
                            "content": delta,
                            "stage": "generating"
                        }, ensure_ascii=False)
            
            # 尝试解析最终的JSON结果
            try:
                # 提取JSON部分
                json_str = self._extract_json(previous_content)
                if json_str:
                    testcase_data = json.loads(json_str)
                    yield json.dumps({
                        "type": "complete",
                        "stage": "finished",
                        "data": testcase_data
                    }, ensure_ascii=False)
                else:
                    yield json.dumps({
                        "type": "complete",
                        "stage": "finished",
                        "content": previous_content
                    }, ensure_ascii=False)
            except json.JSONDecodeError:
                yield json.dumps({
                    "type": "complete",
                    "stage": "finished",
                    "content": previous_content
                }, ensure_ascii=False)
                
        except Exception as e:
            yield json.dumps({
                "type": "error",
                "error": f"测试用例生成失败: {str(e)}"
            }, ensure_ascii=False)

    def _build_vision_prompt(self, context: str, requirements: str) -> str:
        """构建图片分析提示词"""
        prompt_parts = ["请详细分析这张图片的内容，提取测试需求："]
        
        if context:
            prompt_parts.append(f"\n【上下文背景】\n{context}")
        
        if requirements:
            prompt_parts.append(f"\n【用户要求】\n{requirements}")
        
        prompt_parts.append("""
请按以下结构输出分析结果：
1. 功能模块识别
2. 业务流程梳理  
3. 测试需求提取
4. 边界条件和异常场景
5. 测试优先级建议""")
        
        return "\n".join(prompt_parts)

    def _build_testcase_prompt(
        self, 
        analysis_result: str, 
        context: str, 
        requirements: str
    ) -> str:
        """构建测试用例生成提示词"""
        prompt = f"""基于以下需求分析结果，生成完整的测试用例：

【需求分析结果】
{analysis_result}
"""
        
        if context:
            prompt += f"\n【项目背景】\n{context}\n"
        
        if requirements:
            prompt += f"\n【特殊要求】\n{requirements}\n"
        
        prompt += """
请生成测试用例，输出格式为JSON数组：
[
  {
    "id": "TC-001",
    "title": "用例标题",
    "module": "所属模块",
    "precondition": "前置条件",
    "steps": ["步骤1", "步骤2"],
    "expected_result": "预期结果",
    "priority": "P0/P1/P2/P3",
    "test_type": "功能/性能/兼容性/安全性/易用性"
  }
]

要求：
1. 用例编号从TC-001开始连续编号
2. 覆盖所有功能点和业务流程
3. 包含正向和反向测试场景
4. 考虑边界值和异常情况
5. 至少生成20条测试用例"""
        
        return prompt

    def _extract_text(self, event: object) -> Optional[str]:
        """从事件中提取文本内容"""
        if isinstance(event, TextMessage):
            return str(event.content)

        if isinstance(event, TaskResult):
            if not event.messages:
                return None
            last = event.messages[-1]
            return str(getattr(last, "content", "")) or None

        content = getattr(event, "content", None)
        if content is None:
            return None
        return str(content)

    def _extract_json(self, text: str) -> Optional[str]:
        """从文本中提取JSON部分"""
        # 查找JSON数组
        start_idx = text.find("[")
        end_idx = text.rfind("]")
        
        if start_idx != -1 and end_idx != -1 and end_idx > start_idx:
            return text[start_idx:end_idx+1]
        
        # 查找JSON对象
        start_idx = text.find("{")
        end_idx = text.rfind("}")
        
        if start_idx != -1 and end_idx != -1 and end_idx > start_idx:
            return text[start_idx:end_idx+1]
        
        return None

    def export_to_excel(self, testcases: List[Dict[str, Any]], filename: Optional[str] = None) -> BytesIO:
        """
        将测试用例导出为Excel文件
        
        Args:
            testcases: 测试用例列表
            filename: 可选的文件名
            
        Returns:
            Excel文件的BytesIO对象
        """
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"测试用例_{timestamp}.xlsx"
        
        # 创建工作簿
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "测试用例"
        
        # 定义样式
        header_font = Font(name='微软雅黑', size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        cell_font = Font(name='微软雅黑', size=10)
        cell_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        
        border = Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000')
        )
        
        # 设置列宽
        column_widths = {
            'A': 12,  # 用例编号
            'B': 35,  # 用例标题
            'C': 18,  # 所属模块
            'D': 30,  # 前置条件
            'E': 50,  # 测试步骤
            'F': 40,  # 预期结果
            'G': 10,  # 优先级
            'H': 12,  # 测试类型
        }
        
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        # 写入表头
        headers = ["用例编号", "用例标题", "所属模块", "前置条件", "测试步骤", "预期结果", "优先级", "测试类型"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        # 设置行高
        ws.row_dimensions[1].height = 30
        
        # 写入数据
        priority_colors = {
            'P0': 'FF6B6B',  # 红色
            'P1': 'FFA94D',  # 橙色
            'P2': 'FFD43B',  # 黄色
            'P3': '69DB7C',  # 绿色
        }
        
        for row_idx, testcase in enumerate(testcases, 2):
            # 处理测试步骤（数组转为字符串）
            steps = testcase.get("steps", [])
            if isinstance(steps, list):
                steps_text = "\n".join([f"{i+1}. {step}" for i, step in enumerate(steps)])
            else:
                steps_text = str(steps)
            
            row_data = [
                testcase.get("id", f"TC-{row_idx-1:03d}"),
                testcase.get("title", ""),
                testcase.get("module", ""),
                testcase.get("precondition", ""),
                steps_text,
                testcase.get("expected_result", ""),
                testcase.get("priority", "P2"),
                testcase.get("test_type", "功能"),
            ]
            
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = cell_font
                cell.alignment = cell_alignment
                cell.border = border
                
                # 根据优先级设置背景色
                if col_idx == 7:  # 优先级列
                    priority = testcase.get("priority", "P2")
                    color = priority_colors.get(priority, 'FFFFFF')
                    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            
            ws.row_dimensions[row_idx].height = 60
        
        # 冻结首行
        ws.freeze_panes = 'A2'
        
        # 保存到BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output
