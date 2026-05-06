from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from typing import List, Dict, Any
import io
from loguru import logger


class ExcelExporter:
    """Excel导出服务"""
    
    @staticmethod
    def export_test_cases(test_cases: List[Dict[str, Any]]) -> bytes:
        """导出测试用例到Excel"""
        
        wb = Workbook()
        ws = wb.active
        ws.title = "测试用例"
        
        # 定义表头样式
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # 定义表头
        headers = ["用例编号", "用例标题", "前置条件", "测试步骤", "预期结果", "优先级", "用例类型"]
        
        # 写入表头
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 设置列宽
        column_widths = [15, 30, 25, 40, 30, 12, 15]
        for col_num, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col_num)].width = width
        
        # 写入数据
        row_num = 2
        for case in test_cases:
            # 用例编号
            ws.cell(row=row_num, column=1).value = case.get('case_number', '')
            
            # 用例标题
            ws.cell(row=row_num, column=2).value = case.get('title', '')
            
            # 前置条件
            ws.cell(row=row_num, column=3).value = case.get('precondition', '')
            
            # 测试步骤
            steps = case.get('steps', [])
            if isinstance(steps, list):
                steps_text = "\n".join([
                    f"{i+1}. {step.get('step', '')}\n   预期: {step.get('expected_result', '')}"
                    for i, step in enumerate(steps)
                ])
            else:
                steps_text = str(steps)
            ws.cell(row=row_num, column=4).value = steps_text
            
            # 预期结果
            ws.cell(row=row_num, column=5).value = case.get('expected_result', '')
            
            # 优先级
            ws.cell(row=row_num, column=6).value = case.get('priority', '')
            
            # 用例类型
            ws.cell(row=row_num, column=7).value = case.get('case_type', '')
            
            # 设置单元格对齐
            for col in range(1, 8):
                ws.cell(row=row_num, column=col).alignment = Alignment(
                    wrap_text=True,
                    vertical="top"
                )
            
            row_num += 1
        
        # 保存到字节流
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        logger.info(f"成功导出 {len(test_cases)} 个测试用例到Excel")
        return output.getvalue()


excel_exporter = ExcelExporter()
